"""Clasificación supervisada Q -> AB sin modificar el libro original."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import joblib
import numpy as np
from openpyxl import load_workbook
import sklearn
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics import accuracy_score, classification_report, f1_score
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import FeatureUnion, Pipeline
from sklearn.svm import LinearSVC

ROOT = Path(__file__).resolve().parent
ARCHIVO = ROOT / "data" / "Clasificacion de Tickets 2021.xlsm"
MODELO = ROOT / "modelos" / "clasificador_tickets.joblib"
SALIDAS = ROOT / "salidas"


def limpiar_texto(valor):
    """Misma normalización para entrenar, agrupar y predecir."""
    if not isinstance(valor, str):
        return ""
    texto = unicodedata.normalize("NFKD", valor.casefold())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(re.findall(r"[^\W_]+", texto, flags=re.UNICODE))


def normalizar_categoria(valor):
    # Solo espacios y capitalización. No fusiona sinónimos ni errores de captura.
    return " ".join(str(valor).split()).casefold() if valor is not None else ""


def leer_tickets(archivo=ARCHIVO, hoja="Data"):
    """Lee posiciones Excel, no nombres aproximados ni columnas auxiliares."""
    with Path(archivo).open("rb") as entrada:
        libro = load_workbook(entrada, read_only=True, data_only=True)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"No existe la hoja {hoja!r}. Hojas: {libro.sheetnames}")
            datos = libro[hoja]
            if datos.max_column < 28:
                raise ValueError("La hoja debe contener las columnas Q y AB.")
            encabezados = next(datos.iter_rows(min_row=1, max_row=1, max_col=28, values_only=True))
            esperados = {16: "Work Description (Close Details)", 27: "Dispositivo"}
            for indice, esperado in esperados.items():
                if str(encabezados[indice]).strip() != esperado:
                    raise ValueError(f"Columna {indice + 1}: se esperaba {esperado!r}, se encontró {encabezados[indice]!r}.")
            filas = []
            for numero, valores in enumerate(datos.iter_rows(min_row=2, max_col=28, values_only=True), 2):
                if all(v is None for v in valores):
                    continue
                texto, categoria = valores[16], valores[27]
                filas.append({"fila_excel": numero, "ticket": str(valores[0] or "").strip(),
                              "texto_q": texto if isinstance(texto, str) else "",
                              "categoria_ab": categoria if categoria is not None else "",
                              "texto": limpiar_texto(texto), "categoria": normalizar_categoria(categoria)})
            return filas
        finally:
            libro.close()


def preparar_datos(filas):
    validas = [f for f in filas if f["texto"] and f["categoria"]]
    etiquetas = defaultdict(set)
    variantes = defaultdict(Counter)
    for f in validas:
        etiquetas[f["texto"]].add(f["categoria"])
        variantes[f["categoria"]][" ".join(str(f["categoria_ab"]).split())] += 1
    conflictos = {t for t, cats in etiquetas.items() if len(cats) > 1}
    limpias = [dict(f) for f in validas if f["texto"] not in conflictos]
    if not limpias:
        raise ValueError("No hay ejemplos Q/AB utilizables después de eliminar contradicciones.")
    # Componentes conectados: textos idénticos O mismo ticket quedan juntos.
    padres = list(range(len(limpias)))
    def raiz(i):
        while padres[i] != i:
            padres[i] = padres[padres[i]]
            i = padres[i]
        return i
    indices = {}
    for i, f in enumerate(limpias):
        claves = [("texto", f["texto"])]
        if f["ticket"]:
            claves.append(("ticket", f["ticket"]))
        for clave in claves:
            if clave in indices:
                padres[raiz(i)] = raiz(indices[clave])
            else:
                indices[clave] = i
    unicas, vistos = [], set()
    for i, f in enumerate(limpias):
        if f["texto"] not in vistos:
            f["grupo"] = raiz(i)
            unicas.append(f)
            vistos.add(f["texto"])
    auditoria = {"filas_leidas": len(filas),
                 "sin_texto_utilizable_q": sum(not f["texto"] for f in filas),
                 "sin_categoria_ab": sum(not f["categoria"] for f in filas),
                 "pares_q_ab": len(validas), "textos_con_etiquetas_contradictorias": len(conflictos),
                 "filas_contradictorias_excluidas": len(validas) - len(limpias),
                 "repeticiones_texto_categoria_excluidas": len(limpias) - len(unicas),
                 "ejemplos_utilizables": len(unicas),
                 "categorias_utilizables": len({f["categoria"] for f in unicas})}
    nombres = {k: c.most_common(1)[0][0] for k, c in variantes.items()}
    return unicas, [f for f in validas if f["texto"] in conflictos], nombres, auditoria


def crear_pipeline():
    return Pipeline([
        ("tfidf", FeatureUnion([
            ("palabras", TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True, max_features=60000)),
            ("caracteres", TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5),
                                          sublinear_tf=True, max_features=80000)),
        ])),
        ("svm", LinearSVC(C=1.0, dual="auto", max_iter=10000, random_state=42)),
    ])


def dividir_datos(filas, semilla=42):
    grupos = [f["grupo"] for f in filas]
    if len(set(grupos)) < 5 or len({f["categoria"] for f in filas}) < 2:
        raise ValueError("Se requieren al menos cinco grupos independientes y dos categorías.")
    entrenamiento, prueba = next(GroupShuffleSplit(n_splits=1, test_size=.2, random_state=semilla).split(filas, groups=grupos))
    if len({filas[i]["categoria"] for i in entrenamiento}) < 2:
        raise ValueError("El entrenamiento tiene menos de dos categorías; se necesitan más ejemplos.")
    return entrenamiento, prueba


def puntuaciones(pipeline, textos):
    resultado = np.asarray(pipeline.decision_function(textos))
    # LinearSVC devuelve un vector para el caso binario.
    return np.column_stack([-resultado, resultado]) if resultado.ndim == 1 else resultado


def guardar_csv(destino, filas, campos):
    destino = Path(destino)
    if destino.suffix.lower() != ".csv":
        raise ValueError("La salida debe tener extensión .csv; no se reescribe el Excel.")
    destino.parent.mkdir(parents=True, exist_ok=True)
    def seguro(v):
        if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + v
        return v
    with destino.open("w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=campos, extrasaction="ignore")
        escritor.writeheader()
        escritor.writerows({k: seguro(v) for k, v in f.items()} for f in filas)


def entrenar(archivo=ARCHIVO, hoja="Data", modelo=MODELO, salida=SALIDAS, semilla=42):
    archivo, modelo, salida = Path(archivo), Path(modelo), Path(salida)
    if modelo.suffix != ".joblib" or modelo.resolve() == archivo.resolve():
        raise ValueError("El modelo debe ser un archivo .joblib distinto del libro.")
    filas, conflictos, nombres, auditoria = preparar_datos(leer_tickets(archivo, hoja))
    train, test = dividir_datos(filas, semilla)
    x = np.array([f["texto"] for f in filas], dtype=object)
    y = np.array([f["categoria"] for f in filas])
    pipeline = crear_pipeline()
    pipeline.fit(x[train], y[train])
    pred = pipeline.predict(x[test])
    clases = pipeline.classes_
    scores = puntuaciones(pipeline, x[test])
    top = clases[np.argsort(-scores, axis=1)[:, :min(3, len(clases))]]
    aciertos_top3 = [real in opciones for real, opciones in zip(y[test], top)]
    mayoritaria = Counter(y[train]).most_common(1)[0][0]
    metricas = {
        "auditoria": auditoria, "semilla": semilla,
        "evaluacion": "80/20 por grupos de texto normalizado y ticket, sin ajuste con la prueba",
        "ejemplos_entrenamiento": len(train), "ejemplos_prueba": len(test),
        "categorias_entrenamiento": len(clases), "categorias_prueba": len(set(y[test])),
        "categorias_prueba_ausentes_entrenamiento": sorted(set(y[test]) - set(y[train])),
        "filas_prueba_categoria_no_aprendida": sum(v not in set(clases) for v in y[test]),
        "exactitud": float(accuracy_score(y[test], pred)),
        "f1_macro_categorias_prueba": float(f1_score(y[test], pred, labels=sorted(set(y[test])), average="macro", zero_division=0)),
        "f1_ponderado": float(f1_score(y[test], pred, average="weighted", zero_division=0)),
        "exactitud_top3": float(np.mean(aciertos_top3)),
        "exactitud_base_mayoritaria": float(np.mean(y[test] == mayoritaria)),
        "por_categoria": classification_report(y[test], pred, output_dict=True, zero_division=0),
        "sklearn": sklearn.__version__,
    }
    salida.mkdir(parents=True, exist_ok=True)
    (salida / "metricas.json").write_text(json.dumps(metricas, ensure_ascii=False, indent=2), encoding="utf-8")
    guardar_csv(salida / "evaluacion.csv", [
        {**filas[i], "categoria_predicha": nombres[p], "acierto": str(y[i] == p),
         "alternativas": " / ".join(nombres[c] for c in alternativas)}
        for i, p, alternativas in zip(test, pred, top)
    ], ["fila_excel", "ticket", "texto_q", "categoria_ab", "categoria_predicha", "acierto", "alternativas"])
    guardar_csv(salida / "conflictos.csv", conflictos, ["fila_excel", "ticket", "texto_q", "categoria_ab"])
    # Modelo de uso: reajuste con TODOS los ejemplos limpios, después de medir.
    pipeline.fit(x, y)
    conteos = dict(Counter(y))
    paquete = {"version": 1, "pipeline": pipeline, "nombres": nombres, "conteos": conteos,
               "hoja": hoja, "columna_texto": "Q", "columna_categoria": "AB",
               "origen": archivo.name, "sha256_origen": hashlib.sha256(archivo.read_bytes()).hexdigest(),
               "sklearn": sklearn.__version__, "metricas": metricas}
    modelo.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(paquete, modelo, compress=3)
    guardar_csv(salida / "categorias.csv", [
        {"categoria": nombres[c], "clave_normalizada": c, "ejemplos_unicos": n}
        for c, n in sorted(conteos.items(), key=lambda item: (-item[1], item[0]))
    ], ["categoria", "clave_normalizada", "ejemplos_unicos"])
    return paquete


def cargar_modelo(modelo=MODELO):
    # joblib/pickle: cargar únicamente archivos propios o de confianza.
    paquete = joblib.load(modelo)
    if paquete.get("version") != 1:
        raise ValueError("Versión de modelo no compatible.")
    return paquete


def predecir_textos(textos, modelo, lote=256):
    if lote < 1:
        raise ValueError("El tamaño del lote debe ser positivo.")
    originales = list(textos)
    normalizados = [limpiar_texto(t) for t in originales]
    salida = [{"categoria": None, "estado": "sin_texto", "alternativas": []} for _ in originales]
    indices = [i for i, t in enumerate(normalizados) if t]
    pipeline, nombres = modelo["pipeline"], modelo["nombres"]
    for inicio in range(0, len(indices), lote):
        seleccion = indices[inicio:inicio + lote]
        features = pipeline.named_steps["tfidf"].transform([normalizados[i] for i in seleccion])
        scores = np.asarray(pipeline.named_steps["svm"].decision_function(features))
        if scores.ndim == 1:
            scores = np.column_stack([-scores, scores])
        top = np.argsort(-scores, axis=1)[:, :min(3, len(pipeline.classes_))]
        for j, i in enumerate(seleccion):
            if features[j].nnz == 0:
                salida[i]["estado"] = "sin_vocabulario_conocido"
                continue
            alternativas = [{"categoria": nombres[pipeline.classes_[k]], "puntaje_svm": float(scores[j, k])}
                            for k in top[j]]
            clave = pipeline.classes_[top[j, 0]]
            salida[i] = {"categoria": nombres[clave], "estado": "prediccion",
                         "ejemplos_entrenamiento_categoria": modelo["conteos"][clave],
                         "alternativas": alternativas}
    return salida


def clasificar_texto(texto, modelo):
    """Devuelve una categoría de AB o None si no hay texto/vocabulario."""
    return predecir_textos([texto], modelo)[0]["categoria"]


def clasificar_archivo(archivo=ARCHIVO, modelo=MODELO, salida=SALIDAS / "predicciones.csv", hoja="Data"):
    if Path(salida).resolve() == Path(archivo).resolve():
        raise ValueError("La salida debe ser distinta del archivo de entrada.")
    filas = leer_tickets(archivo, hoja)
    paquete = cargar_modelo(modelo)
    resultados = predecir_textos([f["texto_q"] for f in filas], paquete)
    exportacion = []
    for f, resultado in zip(filas, resultados):
        alternativas = resultado["alternativas"]
        exportacion.append({**f, "categoria_predicha": resultado["categoria"] or "",
                            "estado": resultado["estado"],
                            "ejemplos_categoria": resultado.get("ejemplos_entrenamiento_categoria", ""),
                            "alternativa_2": alternativas[1]["categoria"] if len(alternativas) > 1 else "",
                            "alternativa_3": alternativas[2]["categoria"] if len(alternativas) > 2 else ""})
    guardar_csv(salida, exportacion, ["fila_excel", "ticket", "texto_q", "categoria_ab", "categoria_predicha",
                                     "estado", "ejemplos_categoria", "alternativa_2", "alternativa_3"])
    return len(exportacion)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    acciones = parser.add_subparsers(dest="accion", required=True)
    p = acciones.add_parser("entrenar", help="Entrenar y evaluar Q -> AB")
    p.add_argument("--archivo", type=Path, default=ARCHIVO)
    p.add_argument("--hoja", default="Data")
    p.add_argument("--modelo", type=Path, default=MODELO)
    p.add_argument("--salida", type=Path, default=SALIDAS)
    p.add_argument("--semilla", type=int, default=42)
    p = acciones.add_parser("predecir", help="Clasificar un texto")
    p.add_argument("texto")
    p.add_argument("--modelo", type=Path, default=MODELO)
    p = acciones.add_parser("archivo", help="Clasificar toda Q y exportar CSV")
    p.add_argument("--archivo", type=Path, default=ARCHIVO)
    p.add_argument("--hoja", default="Data")
    p.add_argument("--modelo", type=Path, default=MODELO)
    p.add_argument("--salida", type=Path, default=SALIDAS / "predicciones.csv")
    args = vars(parser.parse_args())
    accion = args.pop("accion")
    try:
        if accion == "entrenar":
            paquete = entrenar(**args)
            m = paquete["metricas"]
            print(json.dumps({k: v for k, v in m.items() if k != "por_categoria"}, ensure_ascii=False, indent=2))
        elif accion == "predecir":
            print(json.dumps(predecir_textos([args["texto"]], cargar_modelo(args["modelo"]))[0], ensure_ascii=False, indent=2))
        else:
            print(f"Filas exportadas: {clasificar_archivo(**args)}. Salida: {args['salida']}")
    except (ValueError, OSError, KeyError) as error:
        parser.exit(2, f"Error: {error}\n")


if __name__ == "__main__":
    main()
